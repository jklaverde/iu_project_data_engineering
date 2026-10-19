"""Environmental officer reports (D55, UC-15-UC-18, FR-O2-O9)."""

import re
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import environment as env
from app import report_xlsx
from app import reports as rep

UTC = timezone.utc
CO = {"mean": 0.005, "stddev": 0.001, "ceiling": 0.012}
THRESHOLDS = {
    "co": CO,
    "lpg": {"mean": 0.008, "stddev": 0.001, "ceiling": 0.02},
    "smoke": {"mean": 0.02, "stddev": 0.002, "ceiling": 0.05},
    "temp": {"mean": 22.0, "stddev": 0.5, "ceiling": None},
    "humidity": {"mean": 50.0, "stddev": 2.0, "ceiling": None},
}
META = [
    {"device_id": "a", "name": "Kaiserstrasse", "area": "B70", "lat": 52.5, "lon": 7.3},
    {"device_id": "b", "name": "Baccum", "area": "Residential", "lat": 52.5, "lon": 7.3},
]


def row(start: datetime, *, co=0.005, co_min=None, co_max=None, anomalies=0, events=60, device="a") -> dict:
    """An agg_1h row; every metric other than co sits at its seasonal mean."""
    # The format CassandraReader really returns: naive UTC, no "Z" (the driver
    # hands back naive datetimes). Using "Z" here hid a bug once.
    window_start = start.astimezone(UTC).replace(tzinfo=None).isoformat(timespec="milliseconds")
    r = {"device_id": device, "window_start": window_start, "event_count": events, "anomaly_count": anomalies}
    for metric in ("lpg", "smoke", "temp", "humidity"):
        mean = env.seasonal_mean(metric, THRESHOLDS[metric], start)
        r.update({f"{metric}_avg": mean, f"{metric}_min": mean, f"{metric}_max": mean})
    r.update({"co_avg": co, "co_min": co if co_min is None else co_min, "co_max": co if co_max is None else co_max})
    return r


def hours(start: datetime, n: int, **kw) -> list:
    return [row(start + timedelta(hours=i), **kw) for i in range(n)]


# ------------------------------------------------------------ local time ---

def test_reader_timestamps_without_an_offset_are_read_as_utc():
    assert env.parse_iso("2026-09-01T10:00:00.000") == datetime(2026, 9, 1, 10, tzinfo=UTC)
    assert env.parse_iso("2026-09-01T10:00:00.000Z") == datetime(2026, 9, 1, 10, tzinfo=UTC)


@pytest.mark.parametrize("utc, local_hour", [
    (datetime(2026, 1, 15, 12, tzinfo=UTC), 13),       # winter: +1
    (datetime(2026, 7, 15, 12, tzinfo=UTC), 14),       # summer: +2
    (datetime(2026, 3, 29, 0, 59, tzinfo=UTC), 1),     # just before the spring change
    (datetime(2026, 3, 29, 1, 0, tzinfo=UTC), 3),      # 02:00 local is skipped
    (datetime(2026, 10, 25, 0, 59, tzinfo=UTC), 2),    # last minute of summer time
    (datetime(2026, 10, 25, 1, 0, tzinfo=UTC), 2),     # 02:00 local happens twice
])
def test_berlin_time_follows_the_eu_daylight_saving_rule(utc, local_hour):
    assert rep.to_local(utc).hour == local_hour


def test_month_bounds_are_local_midnights():
    start, end = rep.month_bounds(2026, 3)
    assert start == datetime(2026, 2, 28, 23, tzinfo=UTC)   # 1 March 00:00 CET
    assert end == datetime(2026, 3, 31, 22, tzinfo=UTC)     # 1 April 00:00 CEST
    assert rep.month_bounds(2026, 12)[1] == datetime(2026, 12, 31, 23, tzinfo=UTC)


def test_shift_month_crosses_years():
    assert rep.shift_month(2026, 1, -1) == (2025, 12)
    assert rep.shift_month(2026, 9, -12) == (2025, 9)
    assert rep.shift_month(2025, 12, 1) == (2026, 1)


# ------------------------------------------------------- hour judgement ---

T = datetime(2026, 7, 10, 10, tzinfo=UTC)


def test_an_hour_needs_a_pipeline_flag_to_count_as_out_of_range():
    # Same value as the out-of-range log: a high maximum without an anomaly
    # flag from the pipeline is not counted.
    assert not rep.judge_hour(row(T, co_max=0.011), "co", CO)["out_of_range"]
    assert rep.judge_hour(row(T, co_max=0.011, anomalies=1), "co", CO) == {"out_of_range": True, "over_limit": False}


def test_over_the_ceiling_is_also_over_limit():
    assert rep.judge_hour(row(T, co_max=0.02, anomalies=2), "co", CO) == {"out_of_range": True, "over_limit": True}


def test_a_flagged_hour_inside_the_normal_range_is_not_counted():
    assert not rep.judge_hour(row(T, anomalies=3), "co", CO)["out_of_range"]


def test_below_the_normal_range_counts_too():
    low, _ = rep.hour_band("temp", THRESHOLDS["temp"], T)
    r = row(T, anomalies=1)
    r["temp_min"] = low - 1
    assert rep.judge_hour(r, "temp", THRESHOLDS["temp"])["out_of_range"]


def test_the_hour_band_matches_the_timeline_point_band():
    # The log on screen judges a 1h point by attach_normal_bands; the reports must use the same band.
    point = [{"window_start": rep.iso_utc(T)}]
    env.attach_normal_bands(point, "temp", THRESHOLDS["temp"], "1h")
    assert rep.hour_band("temp", THRESHOLDS["temp"], T) == (point[0]["normal_min"], point[0]["normal_max"])


# -------------------------------------------------------------- summaries ---

def test_metric_summary_counts_and_weights():
    rows = [row(T, co=0.004, events=30), row(T + timedelta(hours=1), co=0.010, co_max=0.02, anomalies=1, events=90)]
    s = rep.metric_summary(rows, "co", CO)
    assert s["hours_with_data"] == 2
    assert s["out_of_range_hours"] == 1 and s["over_limit_hours"] == 1
    assert s["average"] == pytest.approx((0.004 * 30 + 0.010 * 90) / 120)
    assert s["peak"] == 0.02
    assert s["peak_at"] == "2026-07-10T13:00"  # local time


def test_episodes_merge_consecutive_hours_only():
    rows = hours(T, 3, co_max=0.011, anomalies=1) + hours(T + timedelta(hours=5), 1, co_max=0.02, anomalies=1)
    eps = rep.episodes(rows, "a", THRESHOLDS)
    assert [(e["hours"], e["over_limit"]) for e in eps] == [(3, False), (1, True)]
    assert eps[0]["start"] == "2026-07-10T12:00" and eps[0]["end"] == "2026-07-10T15:00"


def test_monthly_report_compares_and_ranks():
    sep_start, _ = rep.month_bounds(2026, 9)
    aug_start, _ = rep.month_bounds(2026, 8)
    ya_start, _ = rep.month_bounds(2025, 9)
    rows = {
        "a": hours(sep_start, 48) + hours(sep_start + timedelta(days=3), 4, co_max=0.011, anomalies=1)
             + hours(aug_start, 24) + hours(aug_start + timedelta(days=2), 1, co_max=0.011, anomalies=1)
             + hours(ya_start, 24, device="a"),
        "b": hours(sep_start, 48, device="b"),
    }
    report = rep.monthly_report(2026, 9, rows, META[::-1], {"a": THRESHOLDS, "b": THRESHOLDS})  # input order != rank
    assert report["month"] == "2026-09" and report["previous_month"] == "2026-08" and report["year_ago_month"] == "2025-09"
    assert [loc["name"] for loc in report["locations"]] == ["Kaiserstrasse", "Baccum"]  # follows the ranking
    top, second = report["ranking"]
    assert (top["name"], top["out_of_range_hours"], top["change_vs_previous"], top["change_vs_year_ago"]) == ("Kaiserstrasse", 4, 3, 4)
    assert (second["name"], second["out_of_range_hours"], second["change_vs_previous"]) == ("Baccum", 0, None)  # no data in August
    assert report["episodes"][0]["hours"] == 4 and report["episodes"][0]["name"] == "Kaiserstrasse"
    co = report["locations"][0]["metrics"]["co"]
    assert (co["current"]["out_of_range_hours"], co["previous"]["out_of_range_hours"]) == (4, 1)


def test_rows_outside_the_month_are_ignored():
    _, end = rep.month_bounds(2026, 9)
    report = rep.monthly_report(2026, 9, {"a": hours(end, 5, co_max=0.02, anomalies=1)}, META[:1], {"a": THRESHOLDS})
    assert report["ranking"][0]["hours_with_data"] == 0


# ---------------------------------------------------------- hourly shapes ---

def test_hour_window_can_cross_midnight():
    night = [h for h in range(24) if rep.in_hour_window(h, 22, 6)]
    assert night == [0, 1, 2, 3, 4, 5, 22, 23]
    assert [h for h in range(24) if rep.in_hour_window(h, 8, 10)] == [8, 9]
    assert all(rep.in_hour_window(h, 5, 5) for h in range(24))


def test_heatmap_is_month_by_local_hour():
    end = datetime(2026, 9, 30, 12, tzinfo=UTC)
    rows = hours(datetime(2026, 9, 1, 22, tzinfo=UTC), 2, co=0.008) + hours(datetime(2026, 8, 5, 10, tzinfo=UTC), 1, co=0.002)
    hm = rep.heatmap(rows, "co", end)
    assert len(hm["months"]) == 12 and hm["months"][-1] == "2026-09" and hm["months"][0] == "2025-10"
    assert hm["values"][-1][0] == pytest.approx(0.008)   # 22:00 UTC = 00:00 local in September
    assert hm["values"][-1][1] == pytest.approx(0.008)
    assert hm["values"][-2][12] == pytest.approx(0.002)  # August, 12:00 local
    assert hm["values"][-1][12] is None
    assert (hm["min"], hm["max"], hm["peak_hour"]) == (pytest.approx(0.002), pytest.approx(0.008), 0)


def test_inquiry_separates_the_hour_window_from_the_rest_of_the_day():
    start, end = rep.local_midnight_utc(date(2026, 9, 1)), rep.local_midnight_utc(date(2026, 9, 3))
    rows = []
    for r in hours(start, 48):
        if rep.in_hour_window(rep.to_local(rep.environment.parse_iso(r["window_start"])).hour, 22, 6):
            r.update(co_avg=0.008, co_max=0.011, anomaly_count=1)
        rows.append(r)
    report = rep.inquiry_report(META[0], THRESHOLDS, rows, rows, start, end, 22, 6)
    co = report["metrics"]["co"]
    assert (co["inside"]["hours_with_data"], co["outside"]["hours_with_data"]) == (16, 32)
    assert (co["inside"]["out_of_range_hours"], co["outside"]["out_of_range_hours"]) == (16, 0)
    assert co["inside"]["average"] > co["outside"]["average"]
    # The first night is 00:00-06:00 (6 h, the night started before the period), then 22:00-06:00 (8 h), then 22:00-24:00.
    assert [e["hours"] for e in report["episodes"]] == [6, 8, 2]
    assert co["profile"][23] == pytest.approx(0.008) and co["profile"][12] == pytest.approx(0.005)


# ------------------------------------------------------------ fact sheet ---

def test_character_labels_pollutants_and_gives_temperature_a_range():
    steady = hours(T, 24, co=0.005)
    spiky = [row(T + timedelta(hours=i), co=0.002 if i % 2 else 0.010) for i in range(24)]
    assert rep.character(steady, "co")["label"] == "steady"
    assert rep.character(spiky, "co")["label"] == "spiky"
    temp = rep.character(steady, "temp")
    assert temp["label"] is None and temp["low"] <= temp["high"]


def test_fact_sheet_uses_the_last_30_days_for_figures_and_12_months_for_the_heatmap():
    now = datetime(2026, 9, 30, 12, tzinfo=UTC)
    old = hours(now - timedelta(days=100), 5, co_max=0.02, anomalies=1)
    recent = hours(now - timedelta(days=2), 24)
    sheet = rep.fact_sheet(META[0], THRESHOLDS, old + recent, now)
    assert sheet["hours_with_data"] == 24
    assert sheet["metrics"]["co"]["out_of_range_hours"] == 0
    daily = sheet["metrics"]["co"]["daily"]
    assert len(daily) == 31  # every local day of the 30-day window, inclusive of today
    assert sum(d["average"] is not None for d in daily) == 2  # 24 hours starting at 14:00 local span two days
    assert daily[0]["average"] is None  # days without data stay on the axis as gaps
    assert sum(v is not None for row_ in sheet["heatmaps"]["co"]["values"] for v in row_) > 0


# --------------------------------------------------- caveats and workbooks ---

INTERNAL = re.compile(r"\b(D\d{1,3}|UC-\d+|FR-[A-Z]+\d*|NFR-\d+)\b|REQUIREMENTS\.md")


def test_caveats_are_present_and_free_of_internal_codes():
    text = " ".join(rep.CAVEATS)
    for needle in ("not legal limit", "derived from a 2020", "simulated"):
        assert needle in text
    for caveat in rep.CAVEATS:
        assert not INTERNAL.search(caveat), caveat


def test_reports_never_include_pressure():
    assert "pressure" not in rep.METRICS


def _sheets(body: bytes) -> dict:
    wb = load_workbook(BytesIO(body))
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def test_monthly_workbook_has_notes_ranking_summary_and_episodes():
    sep_start, _ = rep.month_bounds(2026, 9)
    rows = {"a": hours(sep_start, 6, co_max=0.011, anomalies=1), "b": hours(sep_start, 6, device="b")}
    report = rep.monthly_report(2026, 9, rows, META, {"a": THRESHOLDS, "b": THRESHOLDS})
    sheets = _sheets(report_xlsx.monthly_workbook(report))
    assert list(sheets) == ["Notes", "Ranking", "Summary", "Episodes"]
    notes = " ".join(str(r[0]) for r in sheets["Notes"] if r[0])
    assert all(c in notes for c in rep.CAVEATS) and not INTERNAL.search(notes)
    assert sheets["Ranking"][1][:2] == [1, "Kaiserstrasse"]
    assert len(sheets["Summary"]) == 1 + len(META) * len(rep.METRICS)
    assert sheets["Episodes"][1][4] == 6


def test_extract_workbook_has_hourly_rows_in_local_time_and_the_out_of_range_hours():
    start = datetime(2026, 1, 10, 0, tzinfo=UTC)
    rows = {"a": hours(start, 3) + hours(start + timedelta(hours=3), 2, co_max=0.02, anomalies=1)}
    body = report_xlsx.extract_workbook(META[:1], rows, {"a": THRESHOLDS}, ["co", "temp"], "2026-01-10", "2026-01-11")
    sheets = _sheets(body)
    assert list(sheets) == ["Notes", "Hourly data", "Out-of-range hours"]
    header, first = sheets["Hourly data"][0], sheets["Hourly data"][1]
    assert header[:5] == ["Location", "Device", "Hour (local)", "Hour (UTC)", "Readings"] and len(header) == 5 + 3 * 2
    assert first[2] == datetime(2026, 1, 10, 1) and first[3] == datetime(2026, 1, 10, 0)
    assert len(sheets["Hourly data"]) == 1 + 5
    oor = sheets["Out-of-range hours"]
    assert len(oor) == 1 + 2 and all(r[1] == "CO" and r[9] == "yes" for r in oor[1:])
