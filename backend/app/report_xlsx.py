"""Excel workbooks for the environmental officer's reports (D55, FR-O3, FR-O6, FR-O8).

Built server-side with openpyxl rather than in the browser: the browser-side
options either stopped publishing current versions on npm or pull in a large
transitive tree, and the frontend takes no new npm dependency (NFR-10).

Every workbook opens on a Notes sheet that carries the report's caveats
(FR-O9) - a spreadsheet is forwarded without the app around it.
"""

from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import Cell, WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import environment, reports

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2F5D8A")
_TITLE_FONT = Font(bold=True, size=14)
_DATETIME_FORMAT = "yyyy-mm-dd hh:mm"
_VALUE_FORMAT = "0.0000"


def _row(ws, values: list, formats: dict | None = None) -> None:
    """Appends one row, number formats applied per column (1-based). Formats
    are set on the cells before appending: looking the row up afterwards via
    ws.max_row scans every cell of the sheet, which made a one-year extract
    (26,000 rows) quadratic and the request time out (D55)."""
    formats = formats or {}
    cells = []
    for col, value in enumerate(values, start=1):
        cell = Cell(ws, value=value)
        if col in formats and value is not None:
            cell.number_format = formats[col]
        cells.append(cell)
    ws.append(cells)


def _header(ws, columns: list) -> None:
    ws.append(columns)
    for cell in ws[ws.max_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    ws.freeze_panes = f"A{ws.max_row + 1}"  # a string: ws.cell() would create an empty row


def _widths(ws, minimum: int = 10, maximum: int = 48) -> None:
    for i, column in enumerate(ws.iter_cols(values_only=True), start=1):
        longest = max((len(str(v)) for v in column if v is not None), default=0)
        ws.column_dimensions[get_column_letter(i)].width = max(minimum, min(maximum, longest + 2))


def _notes(wb: Workbook, title: str, lines: list) -> None:
    ws = wb.active
    ws.title = "Notes"
    ws.append([title])
    ws["A1"].font = _TITLE_FONT
    ws.append([f"Generated {reports.to_local(datetime.now(timezone.utc)).strftime('%Y-%m-%d %H:%M')} (Europe/Berlin)"])
    ws.append([])
    for line in lines:
        ws.append([line])
    ws.append([])
    ws.append(["About these numbers"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for caveat in reports.CAVEATS:
        ws.append([caveat])
    ws.append([])
    ws.append(["Units"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for metric in reports.METRICS:
        ws.append([f"{reports.METRIC_LABELS[metric]}: {environment.METRIC_UNITS[metric]}"])
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _local(iso: str | None):
    """reports' local ISO strings / agg_1h UTC strings -> naive local datetimes for Excel."""
    if not iso:
        return None
    if iso.endswith("Z") or "+" in iso[10:]:
        return reports.to_local(environment.parse_iso(iso))
    return datetime.fromisoformat(iso)


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def monthly_workbook(report: dict) -> bytes:
    wb = Workbook()
    _notes(wb, f"Monthly environment report — {report['month']}", [
        f"Compared with {report['previous_month']} (previous month) and {report['year_ago_month']} (same month a year earlier).",
        "Sheets: Ranking (locations by hours out of range), Summary (every location and metric), "
        "Episodes (the longest periods out of range).",
    ])

    ws = wb.create_sheet("Ranking")
    _header(ws, ["Rank", "Location", "Area", "Hours with data", "Hours out of range (any metric)",
                 "Hours over the safety limit", "Change vs previous month", "Change vs a year ago"])
    for i, r in enumerate(report["ranking"], start=1):
        ws.append([i, r["name"], r["area"], r["hours_with_data"], r["out_of_range_hours"],
                   r["over_limit_hours"], r["change_vs_previous"], r["change_vs_year_ago"]])
    _widths(ws)

    ws = wb.create_sheet("Summary")
    _header(ws, ["Location", "Metric", "Unit", "Hours with data", "Hours out of range", "Share out of range",
                 "Hours over the safety limit", "Average", "Peak", "Peak at",
                 "Out of range, previous month", "Out of range, a year ago"])
    for loc in report["locations"]:
        for metric in reports.METRICS:
            m = loc["metrics"][metric]
            cur = m["current"]
            _row(ws, [
                loc["name"], reports.METRIC_LABELS[metric], environment.METRIC_UNITS[metric],
                cur["hours_with_data"], cur["out_of_range_hours"], cur["out_of_range_share"],
                cur["over_limit_hours"] if metric in reports.POLLUTANTS else None,
                cur["average"], cur["peak"], _local(cur["peak_at"]),
                m["previous"]["out_of_range_hours"] if m["previous"]["hours_with_data"] else None,
                m["year_ago"]["out_of_range_hours"] if m["year_ago"]["hours_with_data"] else None,
            ], {6: "0.0%", 8: _VALUE_FORMAT, 9: _VALUE_FORMAT, 10: _DATETIME_FORMAT})
    _widths(ws)

    ws = wb.create_sheet("Episodes")
    _header(ws, ["Location", "Metric", "Start", "End", "Hours", "Peak", "Peak at", "Over the safety limit"])
    for e in report["episodes"]:
        _row(ws, [e["name"], reports.METRIC_LABELS[e["metric"]], _local(e["start"]), _local(e["end"]),
                  e["hours"], e["peak"], _local(e["peak_at"]), "yes" if e["over_limit"] else "no"],
             {3: _DATETIME_FORMAT, 4: _DATETIME_FORMAT, 6: _VALUE_FORMAT, 7: _DATETIME_FORMAT})
    _widths(ws)
    return _save(wb)


def extract_workbook(locations: list, rows_by_device: dict, thresholds: dict, metrics: list,
                     start_local: str, end_local: str) -> bytes:
    """UC-18 / FR-O6. locations: device_metadata entries; rows_by_device: agg_1h rows.

    A write-only workbook: a year of three locations is ~26,000 rows and
    ~670,000 cells, and openpyxl's normal mode keeps a Python object per cell
    (about a minute on the local cluster, D55). Write-only streams the rows,
    so column widths are set up front instead of measured afterwards."""
    wb = Workbook(write_only=True)

    notes = wb.create_sheet("Notes")
    notes.column_dimensions["A"].width = 120
    lines = [
        ("Data extract — hourly values", _TITLE_FONT),
        (f"Generated {reports.to_local(datetime.now(timezone.utc)).strftime('%Y-%m-%d %H:%M')} (Europe/Berlin)", None),
        (None, None),
        (f"Period: {start_local} to {end_local} (local time, Europe/Berlin; end exclusive).", None),
        (f"Locations: {', '.join(loc['name'] for loc in locations)}.", None),
        (f"Metrics: {', '.join(reports.METRIC_LABELS[m] for m in metrics)}.", None),
        ("Hourly data: one row per location and hour, with the hour's average, lowest and highest value per "
         "metric and the number of readings in it. Out-of-range hours: every hour and metric outside its "
         "normal range, with the range that applied.", None),
        (None, None),
        ("About these numbers", Font(bold=True)),
        *[(c, None) for c in reports.CAVEATS],
        (None, None),
        ("Units", Font(bold=True)),
        *[(f"{reports.METRIC_LABELS[m]}: {environment.METRIC_UNITS[m]}", None) for m in reports.METRICS],
    ]
    for text, font in lines:
        cell = WriteOnlyCell(notes, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            cell.font = font
        notes.append([cell])

    def header(ws, columns: list, widths: list) -> None:
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        cells = []
        for title in columns:
            cell = WriteOnlyCell(ws, value=title)
            cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
            cells.append(cell)
        ws.append(cells)

    def formatted(ws, value, fmt):
        cell = WriteOnlyCell(ws, value=value)
        cell.number_format = fmt
        return cell

    ws = wb.create_sheet("Hourly data")
    columns = ["Location", "Device", "Hour (local)", "Hour (UTC)", "Readings"]
    for m in metrics:
        label = f"{reports.METRIC_LABELS[m]} ({environment.METRIC_UNITS[m]})"
        columns += [f"{label} avg", f"{label} min", f"{label} max"]
    header(ws, columns, [16, 20, 17, 17, 10] + [len(c) + 2 for c in columns[5:]])
    for loc in locations:
        for r in rows_by_device.get(loc["device_id"], []):
            utc = environment.parse_iso(r["window_start"])
            ws.append([
                loc["name"], loc["device_id"],
                formatted(ws, reports.to_local(utc), _DATETIME_FORMAT),
                formatted(ws, utc.replace(tzinfo=None), _DATETIME_FORMAT),
                r.get("event_count"),
            ] + [r.get(f"{m}_{k}") for m in metrics for k in ("avg", "min", "max")])

    ws = wb.create_sheet("Out-of-range hours")
    header(ws, ["Location", "Metric", "Unit", "Hour (local)", "Lowest", "Highest",
                "Normal range from", "Normal range to", "Safety limit", "Over the safety limit"],
           [16, 13, 9, 17, 12, 12, 19, 17, 13, 21])
    for loc in locations:
        device_thresholds = thresholds.get(loc["device_id"], {})
        for r, metric, verdict in reports.out_of_range_hours(rows_by_device.get(loc["device_id"], []), device_thresholds):
            if metric not in metrics:
                continue
            stats = device_thresholds[metric]
            start = environment.parse_iso(r["window_start"])
            low, high = reports.hour_band(metric, stats, start)
            ceiling = stats.get("ceiling") if metric in reports.POLLUTANTS else None
            ws.append([loc["name"], reports.METRIC_LABELS[metric], environment.METRIC_UNITS[metric],
                       formatted(ws, reports.to_local(start), _DATETIME_FORMAT),
                       r.get(f"{metric}_min"), r.get(f"{metric}_max"), low, high, ceiling,
                       "yes" if verdict["over_limit"] else "no"])
    return _save(wb)
